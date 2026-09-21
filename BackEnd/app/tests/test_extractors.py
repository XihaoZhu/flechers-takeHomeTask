from unittest.mock import Mock

import pymupdf
from docx import Document as DocxDocument

from app.extractors.docx import extract_docx
from app.extractors.pdf import extract_pdf, should_use_ocr


# region validation tests for should_use_ocr function
def test_should_use_ocr_when_page_has_no_text():

    result = should_use_ocr(text="")

    assert result is True


def test_should_use_ocr_when_page_has_few_usable_characters():

    result = should_use_ocr(text="!@#$%^&*()")

    assert result is True


# endregion


# region extraction tests for document types
def test_extract_docx(tmp_path):

    file_path = tmp_path / "test.docx"

    document = DocxDocument()
    document.add_paragraph("Medical Report")
    document.add_paragraph("            ")
    document.add_paragraph("Patient attended hospital.")
    document.save(file_path)

    blocks = extract_docx(file_path)

    assert len(blocks) == 2
    assert blocks[0].text == "Medical Report"
    assert blocks[0].extraction_method == "docx_text"
    assert blocks[1].page_number is None
    assert blocks[1].text == "Patient attended hospital."


def test_extract_pdf_with_digital_text(tmp_path):
    file_path = tmp_path / "digital.pdf"

    orc_client = Mock()

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "This is a digital PDF page with enough embedded text.",
    )
    document.save(file_path)
    document.close()

    blocks = extract_pdf(file_path, ocr_client=orc_client)

    assert len(blocks) == 1
    assert blocks[0].extraction_method == "pdf_text"
    assert "digital PDF" in blocks[0].text
    orc_client.extract_text.assert_not_called()


def test_extract_pdf_uses_ocr_when_no_embedded_text(tmp_path):
    file_path = tmp_path / "scanned.pdf"

    document = pymupdf.open()
    document.new_page()
    document.save(file_path)
    document.close()

    orc_client = Mock()
    orc_client.extract_text.return_value = "This is the OCR extracted text."

    blocks = extract_pdf(file_path, ocr_client=orc_client)

    assert len(blocks) == 1
    assert blocks[0].extraction_method == "pdf_ocr"
    assert blocks[0].page_number == 1
    assert blocks[0].text == "This is the OCR extracted text."
    orc_client.extract_text.assert_called_once()


# endregion
